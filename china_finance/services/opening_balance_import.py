"""Read-only spreadsheet parsing. All database/file permissions belong to caller."""

import csv
import io
import json
import zipfile
from decimal import Decimal

from china_finance.services.opening_balance_validation import amount

MAX_ROWS = 5000
MAX_BYTES = 2 * 1024 * 1024
HEADERS = {
	"科目": "account", "会计科目": "account", "account": "account",
	"往来类型": "party_type", "party_type": "party_type", "往来单位": "party", "party": "party",
	"科目币种": "account_currency", "币种": "account_currency", "account_currency": "account_currency",
	"汇率": "exchange_rate", "exchange_rate": "exchange_rate",
	"借方": "debit", "借方金额": "debit", "debit": "debit",
	"贷方": "credit", "贷方金额": "credit", "credit": "credit",
	"本位币借方": "base_debit", "base_debit": "base_debit", "本位币贷方": "base_credit", "base_credit": "base_credit",
	"摘要": "remarks", "备注": "remarks", "remarks": "remarks", "会计维度": "dimensions_json", "dimensions_json": "dimensions_json",
}
NUMERIC = {"debit", "credit", "base_debit", "base_credit", "exchange_rate"}


def _error(row, field, message):
	return {"row_id": str(row), "field": field, "message": message}


def parse_import(text=None, content=None, filename=None):
	formulas = set()
	try:
		if content is not None:
			if len(content) > MAX_BYTES:
				raise ValueError("导入文件不能超过 2 MB")
			if (filename or "").lower().endswith(".xlsx"):
				with zipfile.ZipFile(io.BytesIO(content)) as archive:
					if sum(info.file_size for info in archive.infolist()) > 20 * MAX_BYTES:
						raise ValueError("Excel 解压后过大")
				from openpyxl import load_workbook
				workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
				try:
					sheet = workbook.active
					if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > len(HEADERS):
						raise ValueError("导入最多允许 5000 行，且不能包含多余列")
					cells = list(sheet.iter_rows(values_only=False))
					formulas = {(row_index, column_index) for row_index, row in enumerate(cells, 1) for column_index, cell in enumerate(row) if cell.data_type == "f"}
					table = [[cell.value for cell in row] for row in cells]
				finally:
					workbook.close()
			else:
				text = content.decode("utf-8-sig")
				table = None
		else:
			table = None
		if table is None:
			if not isinstance(text, str) or len(text.encode()) > MAX_BYTES:
				raise ValueError("请粘贴表格内容，且不能超过 2 MB")
			reader = csv.reader(io.StringIO(text.lstrip("\ufeff")), delimiter="\t" if "\t" in text.split("\n", 1)[0] else ",")
			table = list(reader)
		if not table or len(table) > MAX_ROWS + 1:
			raise ValueError("表格为空或超过 5000 行")
	except (ValueError, TypeError, UnicodeError, zipfile.BadZipFile, ImportError, OSError) as exc:
		return {"rows": [], "errors": [_error(0, "file", str(exc))]}
	fields = [HEADERS.get(str(value or "").strip()) for value in table[0]]
	errors = []
	for index, field in enumerate(fields):
		if not field:
			errors.append(_error(1, "header", "不支持的列名：" + str(table[0][index] or "")))
		elif fields.count(field) > 1:
			errors.append(_error(1, field, "列名重复"))
	if "account" not in fields or not ({"debit", "credit"} & set(fields)):
		errors.append(_error(1, "header", "必须包含科目和借方或贷方列"))
	if errors:
		return {"rows": [], "errors": errors}
	rows = []
	for number, cells in enumerate(table[1:], 2):
		if not any(value not in (None, "") for value in cells):
			continue
		if len(cells) > len(fields) and any(value not in (None, "") for value in cells[len(fields):]):
			errors.append(_error(number, "row", "列数超过表头"))
		row = {"row_id": "import-" + str(number)}
		for column, (field, value) in enumerate(zip(fields, cells)):
			if (number, column) in formulas:
				errors.append(_error(number, field, "Excel 公式必须先转换为明确的值"))
			value = str(value) if value is not None else ""
			if field != "remarks":
				value = value.strip()
			if field in NUMERIC and value:
				try:
					amount(value, precision=9)
					if value.startswith(("=", "+", "@")):
						raise ValueError("金额不能包含公式")
				except ValueError as exc:
					errors.append(_error(number, field, str(exc)))
			if field == "dimensions_json" and value:
				try:
					if not isinstance(json.loads(value), dict):
						raise ValueError("会计维度必须为 JSON 对象")
				except (ValueError, TypeError) as exc:
					errors.append(_error(number, field, str(exc)))
			row[field] = value
		if not row.get("account"):
			errors.append(_error(number, "account", "科目不能为空"))
		rows.append(row)
	return {"rows": rows, "errors": errors}


def download_template_bytes():
	from openpyxl import Workbook

	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "期初余额"
	sheet.append(["科目", "往来类型", "往来单位", "科目币种", "汇率", "借方", "贷方", "本位币借方", "本位币贷方", "摘要", "会计维度"])
	sheet.freeze_panes = "A2"
	for column in sheet.columns:
		sheet.column_dimensions[column[0].column_letter].width = 20
	output = io.BytesIO()
	workbook.save(output)
	return output.getvalue()
