"""Opening balances: exact arithmetic/evidence, plus permission and transaction gates.

Pure checks run without a Frappe site; native lifecycle coverage is exercised by
integration tests when the suite is run inside bench.
"""

import importlib
import json
import unittest
from decimal import Decimal


def core(test):
	try:
		return importlib.import_module("china_finance.services.opening_balance_validation")
	except ModuleNotFoundError:
		test.fail("Opening-balance exact validation and reconciliation helper is missing")


class TestOpeningArithmetic(unittest.TestCase):
	def rows(self):
		return [
			{"row_id": "one", "account": "Bank", "account_currency": "CNY", "debit": "0.10", "credit": "0", "base_debit": "0.10", "base_credit": "0", "exchange_rate": "1", "remarks": "保留原摘要 未建档往来"},
			{"row_id": "two", "account": "Equity", "account_currency": "CNY", "debit": "0", "credit": "0.10", "base_debit": "0", "base_credit": "0.10", "exchange_rate": "1"},
		]

	def test_decimal_totals_are_exact_and_no_balancing_row_is_invented(self):
		module = core(self)
		rows = self.rows() + [dict(row, row_id=row["row_id"] + "2") for row in self.rows()]
		self.assertEqual(module.totals(rows), {"debit": "0.20", "credit": "0.20", "difference": "0.00"})
		self.assertEqual(len(rows), 4)
		rows[0]["base_debit"] = "0.11"
		self.assertEqual(module.totals(rows)["difference"], "0.01")

	def test_rejects_nonfinite_negative_both_sides_or_excess_precision(self):
		module = core(self)
		for value in ("NaN", "Infinity", "-0.01", "0.001", "1e999", "1000000000000"):
			with self.subTest(value=value), self.assertRaises(ValueError):
				module.amount(value, precision=2)
		with self.assertRaises(ValueError):
			module.amount("999999999999.123456789", precision=9)
		self.assertEqual(module.amount("0.123456789", precision=9), Decimal("0.123456789"))
		for changes in ({"credit": "1"}, {"exchange_rate": "0"}, {"base_debit": "0.11"}):
			row = dict(self.rows()[0], **changes)
			self.assertTrue(module.validate_amounts(row, "CNY"))

	def test_currency_conversion_is_explicit_and_exact_to_base_precision(self):
		module = core(self)
		row = dict(self.rows()[0], account_currency="USD", debit="10.01", exchange_rate="7.123456", base_debit="71.31")
		self.assertEqual(module.validate_amounts(row, "CNY"), [])
		row["base_debit"] = "71.30"
		self.assertTrue(module.validate_amounts(row, "CNY"))

	def test_delta_uses_full_target_net_and_preserves_dimensions_and_parties(self):
		module = core(self)
		previous = self.rows()
		previous[0].update(party_type="Customer", party="Customer A", dimensions_json='{"department":"West"}')
		target = [dict(row) for row in previous]
		target[0].update(debit="0.30", base_debit="0.30", remarks="调整目标摘要")
		target[1].update(credit="0.30", base_credit="0.30")
		delta = module.delta_rows(target, previous)
		self.assertEqual(module.totals(delta), {"debit": "0.20", "credit": "0.20", "difference": "0.00"})
		bank = next(row for row in delta if row["account"] == "Bank")
		self.assertEqual(bank["party"], "Customer A")
		self.assertEqual(bank["remarks"], "调整目标摘要")
		self.assertEqual(json.loads(bank["dimensions_json"]), {"department": "West"})
		self.assertEqual(module.delta_rows(target, target), [])

	def test_reduced_or_removed_balance_generates_opposite_delta_and_rate_is_derived(self):
		module = core(self)
		previous = self.rows()
		target = [dict(row, debit="0", credit="0", base_debit="0", base_credit="0") for row in previous]
		delta = module.delta_rows(target, previous)
		self.assertEqual(next(row for row in delta if row["account"] == "Bank")["credit"], "0.10")
		self.assertEqual(module.totals(delta)["difference"], "0.00")

	def test_delta_rejects_rate_only_difference_with_no_foreign_currency_amount(self):
		module = core(self)
		previous = [dict(self.rows()[0], account_currency="USD", debit="1", base_debit="7", exchange_rate="7")]
		target = [dict(previous[0], base_debit="8", exchange_rate="8")]
		with self.assertRaises(ValueError):
			module.delta_rows(target, previous)

	def test_evidence_compares_gross_columns_and_every_accounting_dimension(self):
		module = core(self)
		rows = self.rows()
		self.assertEqual(module.accounting_totals(rows), module.accounting_totals([dict(row) for row in rows]))
		changed = [dict(row) for row in rows]
		changed[0]["dimensions_json"] = '{"department":"South"}'
		self.assertNotEqual(module.accounting_totals(rows), module.accounting_totals(changed))
		changed = [dict(row, debit="0.20", credit="0.10", base_debit="0.20", base_credit="0.10") for row in rows]
		self.assertNotEqual(module.accounting_totals(rows), module.accounting_totals(changed))

	def test_takeover_preserves_each_historical_amount_remark_pair_not_just_totals(self):
		module = core(self)
		if not hasattr(module, "historical_rows_match"):
			self.fail("Historical row amount and remark evidence is missing")
		original = [dict(self.rows()[0], debit="10", base_debit="10", remarks="原摘要A"), dict(self.rows()[0], debit="20", base_debit="20", remarks="原摘要B")]
		self.assertTrue(module.historical_rows_match(list(reversed(original)), original))
		changed = [dict(original[0], debit="20", base_debit="20"), dict(original[1], debit="10", base_debit="10")]
		self.assertFalse(module.historical_rows_match(changed, original))

	def test_fingerprint_binds_saved_version_sources_and_verbatim_remarks(self):
		module = core(self)
		data = {"company": "Company", "version": 2, "rows": self.rows(), "sources": [{"name": "JE", "hash": "abc"}]}
		first = module.fingerprint(data)
		self.assertEqual(first, module.fingerprint(json.loads(json.dumps(data))))
		data["rows"][0]["remarks"] += " changed"
		self.assertNotEqual(first, module.fingerprint(data))


class TestOpeningImports(unittest.TestCase):
	def parser(self):
		try:
			return importlib.import_module("china_finance.services.opening_balance_import")
		except ModuleNotFoundError:
			self.fail("Opening-balance read-only import parser is missing")

	def test_tsv_exact_headers_and_no_write_with_verbatim_remark(self):
		result = self.parser().parse_import(text="科目\t借方\t贷方\t摘要\nBank\t0.10\t0\t未建档往来 原摘要\nEquity\t0\t0.10\t期初")
		self.assertEqual(result["errors"], [])
		self.assertEqual(result["rows"][0]["debit"], "0.10")
		self.assertEqual(result["rows"][0]["remarks"], "未建档往来 原摘要")

	def test_bad_headers_formulas_and_numeric_values_are_visible_errors(self):
		for text in ("科目\t未知字段\nBank\t1", "account,account,debit\nBank,Bank,1", "account,debit\nBank,=1+1", "account,debit\nBank,NaN"):
			with self.subTest(text=text):
				self.assertTrue(self.parser().parse_import(text=text)["errors"])



	def test_xlsx_template_round_trip_and_formula_cells_are_rejected(self):
		import io

		from openpyxl import Workbook, load_workbook

		parser = self.parser()
		template = parser.download_template_bytes()
		workbook = load_workbook(io.BytesIO(template))
		self.assertEqual(workbook.active.cell(1, 1).value, "科目")
		workbook.active.append(["Bank", "", "", "CNY", 1, 10, 0, 10, 0, "保留原摘要", "{}"])
		data = io.BytesIO()
		workbook.save(data)
		self.assertEqual(parser.parse_import(content=data.getvalue(), filename="opening.xlsx")["errors"], [])
		workbook.active.cell(2, 10).value = '=CONCAT("not trusted")'
		data = io.BytesIO()
		workbook.save(data)
		self.assertTrue(parser.parse_import(content=data.getvalue(), filename="opening.xlsx")["errors"])


class TestOpeningEvidence(unittest.TestCase):
	def source(self):
		return {"name": "JE-1", "company": "Company", "posting_date": "2026-01-01", "docstatus": 1, "voucher_type": "Opening Entry", "is_opening": "Yes"}

	def gl(self):
		return [{"name": "GL-1", "voucher_type": "Journal Entry", "voucher_no": "JE-1", "company": "Company", "posting_date": "2026-01-01", "is_opening": "Yes", "is_cancelled": 0}]

	def test_trust_requires_native_header_all_gl_flags_company_date_and_snapshot(self):
		module = core(self)
		if not hasattr(module, "validate_source_identity"):
			self.fail("Historical opening identity reconciliation is missing")
		self.assertEqual(module.validate_source_identity(self.source(), self.gl(), "Company", "2026-01-01"), [])
		self.assertEqual(module.validate_source_identity(dict(self.source(), voucher_type="Journal Entry"), self.gl(), "Company", "2026-01-01"), [])
		for field, value in (("is_opening", "No"), ("docstatus", 0), ("company", "Other"), ("posting_date", "2026-01-02")):
			with self.subTest(field=field):
				self.assertTrue(module.validate_source_identity(dict(self.source(), **{field: value}), self.gl(), "Company", "2026-01-01"))
		for field, value in (("is_opening", "No"), ("is_cancelled", 1), ("company", "Other"), ("posting_date", "2026-01-02"), ("voucher_no", "Other")):
			with self.subTest(field=field):
				self.assertTrue(module.validate_source_identity(self.source(), [dict(self.gl()[0], **{field: value})], "Company", "2026-01-01"))
		self.assertTrue(module.validate_source_identity(self.source(), [], "Company", "2026-01-01"))


class TestOpeningServiceGates(unittest.TestCase):
	@classmethod
	def setUpClass(cls):
		import importlib.util
		import sys
		from pathlib import Path
		from types import ModuleType
		from unittest.mock import MagicMock, patch

		path = Path(__file__).parents[1] / "services/opening_balance.py"
		if not path.exists():
			cls.service = None
			return
		try:
			import frappe
			cls.service = importlib.import_module("china_finance.services.opening_balance")
		except ModuleNotFoundError:
			fake = ModuleType("frappe")
			fake._ = lambda value: value
			fake.whitelist = lambda *args, **kwargs: lambda method: method
			utils = ModuleType("frappe.utils")
			utils.getdate = lambda value=None: str(value or "2026-01-01")
			utils.now_datetime = lambda: "2026-01-01 00:00:00"
			utils.cint = lambda value: int(value or 0)
			spec = importlib.util.spec_from_file_location("_opening_balance_unit", path)
			module = importlib.util.module_from_spec(spec)
			with patch.dict(sys.modules, {"frappe": fake, "frappe.utils": utils}):
				spec.loader.exec_module(module)
			cls.service = module

	def setUp(self):
		from unittest.mock import MagicMock, patch

		if self.service is None:
			self.fail("Opening balance permission/idempotency/transaction service is missing")
		self.frappe = MagicMock()
		self.frappe.PermissionError = PermissionError
		self.frappe.TimestampMismatchError = RuntimeError
		self.frappe.QueryDeadlockError = type("QueryDeadlockError", (Exception,), {})
		self.frappe.throw.side_effect = lambda message, exc=ValueError: (_ for _ in ()).throw(exc(message))
		patcher = patch.object(self.service, "frappe", self.frappe)
		patcher.start()
		self.addCleanup(patcher.stop)

	def test_saved_version_is_required_and_stale_version_rejected(self):
		for supplied in (None, "", 0, 2):
			with self.subTest(supplied=supplied), self.assertRaises((ValueError, RuntimeError)):
				self.service._check_version({"version": 3}, supplied)
		self.service._check_version({"version": 3}, 3)

	def test_direct_batch_document_writes_and_unsupported_native_prep_are_rejected(self):
		with self.assertRaises(PermissionError):
			self.service.protect_opening_document()
		with self.service.internal_edit():
			self.service.protect_opening_document()
		self.assertFalse(self.service.validate_native_opening_preparation({"name": "forged", "custom_china_opening_batch": "fake"}))

	def test_linked_active_opening_cannot_be_cancelled_or_directly_updated(self):
		from unittest.mock import patch
		with patch.object(self.service, "_linked_active_batch", return_value="OPEN-1"):
			for method in ("before_cancel", "before_update_after_submit"):
				with self.subTest(method=method), self.assertRaises(ValueError):
					self.service.guard_linked_opening_source({"name": "JE-1", "doctype": "Journal Entry"}, method)

	def test_activation_failure_rolls_back_entire_savepoint_and_clears_native_context(self):
		from unittest.mock import MagicMock, patch
		doc = {"name": "OPEN-1", "company": "Company", "status": "Draft"}
		with (
			patch.object(self.service, "_get_batch", return_value=doc),
			patch.object(self.service, "_activate_locked", side_effect=ValueError("submit denied")),
		):
			with self.assertRaisesRegex(ValueError, "submit denied"):
				self.service.activate_opening("OPEN-1", "fingerprint")
		self.frappe.db.savepoint.assert_called_once()
		self.frappe.db.rollback.assert_called_once()
		self.assertFalse(self.service.validate_native_opening_preparation({"name": "JE-1"}))

	def test_active_replay_revalidates_saved_evidence_before_returning_receipt(self):
		from unittest.mock import patch
		doc = {"name": "OPEN-1", "company": "Company", "status": "Active", "activation_fingerprint": "accepted", "receipt_json": '{"source_names":["JE-1"]}'}
		with (
			patch.object(self.service, "_get_batch", return_value=doc),
			patch.object(self.service, "_revalidate_receipt", side_effect=ValueError("source cancelled")) as revalidate,
		):
			with self.assertRaisesRegex(ValueError, "source cancelled"):
				self.service.activate_opening("OPEN-1", "accepted")
		revalidate.assert_called_once_with(doc, lock=True)

	def test_activation_uses_saved_preview_and_never_accepts_stale_fingerprint(self):
		from unittest.mock import patch
		doc = {"name": "OPEN-1", "company": "Company", "status": "Draft"}
		with patch.object(self.service, "_preview", return_value={"fingerprint": "current", "errors": []}):
			with self.assertRaises((ValueError, RuntimeError)):
				self.service._activate_locked(doc, "stale")
		self.frappe.get_doc.assert_not_called()

	def test_native_source_guard_serializes_company_before_current_batch_read(self):
		from unittest.mock import patch

		service = self.service
		if service is None:
			self.fail("Opening service missing")
		events = []
		with patch.object(service, "_lock_company", side_effect=lambda company: events.append("company")), patch.object(service, "_linked_active_batch", side_effect=lambda *args, **kwargs: events.append("batch") or None):
			service.guard_linked_opening_source({"doctype": "Journal Entry", "name": "JE", "company": "Company"}, "before_cancel")
		self.assertEqual(events, ["company", "batch"])

	def test_native_output_must_reconcile_to_explicit_target_before_marking_active(self):
		from unittest.mock import patch

		class Doc(dict):
			__getattr__ = dict.get
			__setattr__ = dict.__setitem__

		doc = Doc(name="OPEN-1", company="Company", opening_date="2026-01-01", mode="New", status="Draft", version=1, revision=1, rows=TestOpeningArithmetic().rows(), source_names_json="[]")
		changed = [dict(row) for row in doc.rows]
		changed[0]["base_debit"] = "0.11"
		with (
			patch.object(self.service, "_preview", return_value={"fingerprint": "current", "errors": []}),
			patch.object(self.service, "_post_native", return_value="JE-1"),
			patch.object(self.service, "_source_evidence", return_value={"name": "JE-1", "hash": "hash", "rows": changed}),
			patch.object(self.service, "_save") as save,
		):
			with self.assertRaises(ValueError):
				self.service._activate_locked(doc, "current")
			save.assert_not_called()

	def test_native_company_cost_center_default_is_saved_before_preview(self):
		from types import SimpleNamespace
		from unittest.mock import patch

		company = SimpleNamespace(name="Company", default_currency="CNY", cost_center="Main - C")
		with patch.object(self.service, "_dimension_definitions", return_value={"cost_center": object()}), patch.object(self.service, "_precision", return_value=2):
			rows = self.service._normalise_rows(TestOpeningArithmetic().rows(), company)
			self.assertEqual(json.loads(rows[0]["dimensions_json"]), {"cost_center": "Main - C"})
			inherited = self.service._normalise_rows(TestOpeningArithmetic().rows(), company, apply_defaults=False)
			self.assertEqual(json.loads(inherited[0]["dimensions_json"]), {})

	def test_deadlock_keeps_original_database_rollback_error(self):
		from unittest.mock import patch

		error = self.frappe.QueryDeadlockError("transaction already rolled back")
		with patch.object(self.service, "_get_batch", return_value={"status": "Draft"}), patch.object(self.service, "_activate_locked", side_effect=error):
			with self.assertRaises(self.frappe.QueryDeadlockError) as raised:
				self.service.activate_opening("OPEN", "fingerprint")
		self.assertIs(raised.exception, error)
		self.frappe.db.rollback.assert_not_called()

	def test_preview_fingerprint_binds_master_and_settings_version(self):
		from unittest.mock import patch

		class Doc(dict):
			__getattr__ = dict.get

		doc = Doc(name="OPEN", company="Company", currency="CNY", opening_date="2026-01-01", status="Draft", version=1, revision=1, mode="Takeover", rows=TestOpeningArithmetic().rows(), source_names_json='["JE-1"]')
		if not hasattr(self.service, "_master_evidence"):
			self.fail("Preview master/settings evidence is missing")
		with (
			patch.object(self.service, "_validation_errors", return_value=[]),
			patch.object(self.service, "_opening_evidence", return_value=[]),
			patch.object(self.service, "_master_evidence", side_effect=[{"version": "one"}, {"version": "two"}]),
			patch.object(self.service, "_precision", return_value=2),
		):
			first = self.service._preview(doc)
			second = self.service._preview(doc)
		self.assertNotEqual(first["fingerprint"], second["fingerprint"])

	def test_all_activation_modes_require_native_create_and_submit_permissions(self):
		from unittest.mock import patch

		self.frappe.has_permission.return_value = False
		for mode in ("New", "Takeover", "Adjustment"):
			with self.subTest(mode=mode), patch.object(self.service, "_get_batch", return_value={"mode": mode}), patch.object(self.service, "_activate_locked", return_value={}), patch.object(self.service, "_envelope", return_value={}):
				with self.assertRaises(PermissionError):
					self.service.activate_opening("OPEN", "hash")

	def test_only_service_batch_save_defers_links_and_never_permissions(self):
		from types import SimpleNamespace

		class Doc:
			doctype = "China Opening Balance"
			def __init__(self):
				self.flags = SimpleNamespace(ignore_links=False)
				self.observed = None
			def insert(self):
				self.observed = self.flags.ignore_links
				return self

		doc = Doc()
		self.service._save(doc, insert=True)
		self.assertTrue(doc.observed)
		self.assertFalse(doc.flags.ignore_links)
		self.assertFalse(getattr(doc.flags, "ignore_permissions", False))

	def test_native_posting_requires_enabled_finance_and_a_verified_technical_snapshot(self):
		import sys
		from types import SimpleNamespace
		from unittest.mock import MagicMock, patch

		if not hasattr(self.service, "_ensure_native_snapshot"):
			self.fail("New opening postings must retain a technical accounting snapshot")
		voucher = MagicMock()
		journal = SimpleNamespace(name="JE-1", company="Company")
		with patch.dict(sys.modules, {"china_finance.services.voucher": voucher}):
			voucher.get_company_settings.return_value = None
			with self.assertRaises(ValueError):
				self.service._ensure_native_snapshot(journal)
			voucher.get_company_settings.return_value = {"enabled": 1}
			voucher.create_voucher_from_source.return_value = None
			with self.assertRaises(ValueError):
				self.service._ensure_native_snapshot(journal)
			voucher.create_voucher_from_source.return_value = "CAV-1"
			self.frappe.db.exists.return_value = True
			self.service._ensure_native_snapshot(journal)
			voucher.create_voucher_from_source.assert_called_with(journal, force=True)

	def test_same_named_child_content_changes_master_fingerprint(self):
		if not hasattr(self.service, "_semantic_master"):
			self.fail("Master evidence must recursively serialize child documents")

		class Child:
			def __init__(self, **values):
				self.values = {"doctype": "Closed Document", "name": "same-child", **values}
			def as_dict(self):
				return self.values
			def __str__(self):
				return "Closed Document: same-child"

		for field, before, after in (("closed", 0, 1), ("allowed", "Accounts User", "System Manager"), ("condition", "old", "new")):
			with self.subTest(field=field):
				one = {"children": [Child(**{field: before})]}
				two = {"children": [Child(**{field: after})]}
				self.assertNotEqual(core(self).fingerprint(self.service._semantic_master(one)), core(self).fingerprint(self.service._semantic_master(two)))

	def test_locked_preview_passes_current_read_to_validation_and_master_evidence(self):
		from unittest.mock import patch

		doc = {"name": "OPEN", "company": "Company", "currency": "CNY", "opening_date": "2026-01-01", "status": "Draft", "version": 1, "revision": 1, "mode": "Takeover", "rows": TestOpeningArithmetic().rows(), "source_names_json": '["JE-1"]'}
		with patch.object(self.service, "_validation_errors", return_value=[]) as validation, patch.object(self.service, "_opening_evidence", return_value=[]), patch.object(self.service, "_master_evidence", return_value={}) as master, patch.object(self.service, "_precision", return_value=2):
			self.service._preview(doc, lock=True)
		validation.assert_called_once_with(doc, lock=True)
		master.assert_called_once_with(doc, lock=True)

	def test_activation_master_documents_children_lists_and_singles_use_current_reads(self):
		from unittest.mock import patch

		doc = {"company": "Company", "currency": "CNY", "rows": TestOpeningArithmetic().rows()}
		self.frappe.db.get_value.return_value = "exists"
		self.frappe.db.exists.return_value = True
		self.frappe.db.get_singles_dict.return_value = {}
		self.frappe.get_doc.return_value = {}
		self.frappe.get_all.return_value = []
		self.frappe.db.get_values.return_value = []
		with patch.object(self.service, "_dimension_definitions", return_value={}):
			try:
				self.service._master_evidence(doc, lock=True)
			except TypeError:
				self.fail("Master evidence does not accept current-read lock mode")
		self.frappe.get_single.assert_not_called()
		self.assertTrue(self.frappe.get_doc.call_args_list)
		self.assertTrue(all(call.kwargs.get("for_update") for call in self.frappe.get_doc.call_args_list))
		self.frappe.get_all.assert_not_called()
		self.assertTrue(self.frappe.db.get_values.call_args_list)
		self.assertTrue(all(call.kwargs.get("for_update") for call in self.frappe.db.get_values.call_args_list))
		self.assertTrue(all(call.kwargs.get("order_by", "").endswith(" asc") for call in self.frappe.db.get_values.call_args_list), "Bare order_by=name reverses ASC/DESC between get_all compatibility and get_values")
		self.assertTrue(all(call.kwargs.get("filters") is not None for call in self.frappe.db.get_values.call_args_list), "Frappe get_values(None) returns None for ordinary DocTypes; unfiltered scans require {}")
		self.assertTrue(all(call.kwargs.get("for_update") for call in self.frappe.db.get_singles_dict.call_args_list))
		self.assertEqual(self.frappe.db.get_singles_dict.call_count, 2)

	def test_envelopes_expose_native_currency_precision(self):
		from types import SimpleNamespace
		from unittest.mock import patch

		self.frappe.get_list.return_value = []
		with patch.object(self.service, "_company_context", return_value=SimpleNamespace(default_currency="CNY")), patch.object(self.service, "_dimension_definitions", return_value={}), patch.object(self.service, "_precision", return_value=6):
			self.assertEqual(self.service._envelope(company="Company").get("precision"), 6)

	def test_only_trusted_opening_submit_uses_atomic_native_path(self):
		import importlib.util
		import sys
		from pathlib import Path
		from types import ModuleType
		from unittest.mock import patch

		class Native:
			def submit(self):
				return "native-default"
			def _submit(self):
				return "native-atomic"

		path = "erpnext.accounts.doctype.journal_entry.journal_entry"
		modules = {".".join(path.split(".")[:index]): ModuleType(".".join(path.split(".")[:index])) for index in range(1, 6)}
		modules[path].JournalEntry = Native
		modules["china_finance.services.opening_balance"] = self.service
		file = Path(__file__).parents[1] / "overrides/journal_entry.py"
		spec = importlib.util.spec_from_file_location("_opening_journal_override_unit", file)
		module = importlib.util.module_from_spec(spec)
		with patch.dict(sys.modules, modules):
			spec.loader.exec_module(module)
			with patch.object(self.service, "validate_native_opening_preparation", return_value=False):
				self.assertEqual(module.ChinaFinanceJournalEntry().submit(), "native-default")
			with patch.object(self.service, "validate_native_opening_preparation", return_value=True):
				self.assertEqual(module.ChinaFinanceJournalEntry().submit(), "native-atomic")


class TestOpeningSchema(unittest.TestCase):
	def test_one_company_base_unique_and_all_documents_service_protected(self):
		from pathlib import Path

		root = Path(__file__).parents[1] / "china_finance/doctype"
		batch = root / "china_opening_balance/china_opening_balance.json"
		self.assertTrue(batch.exists(), "Company opening batch DocType is missing")
		definition = json.loads(batch.read_text())
		fields = {field["fieldname"]: field for field in definition["fields"]}
		self.assertTrue(fields["initialization_key"]["unique"])
		self.assertTrue(fields["adjustment_key"]["unique"])
		self.assertEqual(fields["rows"]["options"], "China Opening Balance Line")
		self.assertEqual(fields["status"]["options"], "Draft\nActive")
		controller = batch.with_suffix(".py").read_text()
		for hook in ("validate", "on_trash", "before_rename"):
			self.assertIn("def " + hook, controller)
		self.assertIn("protect_opening_document", controller)
		line = root / "china_opening_balance_line/china_opening_balance_line.json"
		self.assertTrue(line.exists(), "Company opening line DocType is missing")
		definition = json.loads(line.read_text())
		self.assertEqual(definition["istable"], 1)
		self.assertIn("dimensions_json", {field["fieldname"] for field in definition["fields"]})
		self.assertIn("protect_opening_document", line.with_suffix(".py").read_text())


try:
	import frappe as _native_frappe
except ModuleNotFoundError:
	_native_frappe = None


@unittest.skipUnless(_native_frappe, "Requires an isolated native Frappe test site")
class TestOpeningNative(unittest.TestCase):
	"""Native lifecycle on a new company; roll back every case, never commit."""

	def setUp(self):
		from types import MethodType

		from china_finance.services import opening_balance
		from china_finance.setup.china_coa_profile import get_account_by_number
		from china_finance.tests.test_inline_voucher_account_edit import TestInlineVoucherAccountEdit

		self.service = opening_balance
		self.fixture = TestInlineVoucherAccountEdit
		self._make_company = MethodType(self.fixture._make_company, self)
		self._account = MethodType(self.fixture._account, self)
		self.fixture.setUp(self)
		self.equity = get_account_by_number(self.company, "4001").name
		self.opening_date = "2026-01-01"

	def tearDown(self):
		if hasattr(self, "point"):
			self.fixture.tearDown(self)

	def rows(self, amount="100.00"):
		return [
			{"row_id": "bank", "account": self.bank, "account_currency": "CNY", "exchange_rate": "1", "debit": amount, "credit": "0", "base_debit": amount, "base_credit": "0", "remarks": "原始摘要 未建档往来名称"},
			{"row_id": "equity", "account": self.equity, "account_currency": "CNY", "exchange_rate": "1", "debit": "0", "credit": amount, "base_debit": "0", "base_credit": amount, "remarks": "真实资本来源"},
		]

	def save(self, rows=None, **overrides):
		payload = {"company": self.company, "opening_date": self.opening_date, "mode": "New", "rows": rows if rows is not None else self.rows()}
		payload.update(overrides)
		return self.service.save_opening(json.dumps(payload, ensure_ascii=False))

	def activate(self, batch):
		preview = self.service.preview_opening(batch["name"])
		self.assertEqual(preview["errors"], [], preview)
		return self.service.activate_opening(batch["name"], preview["fingerprint"])

	def counts(self):
		return tuple(_native_frappe.db.count(doctype, {"company": self.company}) for doctype in ("Journal Entry", "GL Entry", "China Accounting Voucher"))

	def historical(self):
		from china_finance.services import voucher

		journal = _native_frappe.get_doc({
			"doctype": "Journal Entry", "company": self.company, "posting_date": self.opening_date,
			"voucher_type": "Opening Entry", "is_opening": "Yes", "user_remark": "历史期初来源",
			"accounts": [{"account": row["account"], "debit_in_account_currency": row["debit"], "credit_in_account_currency": row["credit"], "user_remark": row["remarks"]} for row in self.rows()],
		}).insert()
		self.fixture._mark_ready(self, journal)
		journal = voucher._complete_voucher_workflow(journal)
		journal.reload()
		return journal

	def test_saved_draft_never_creates_native_accounting_and_duplicate_base_is_rejected(self):
		before = self.counts()
		batch = self.save()["batch"]
		self.assertEqual(self.counts(), before)
		self.assertEqual(batch["status"], "Draft")
		self.assertEqual(batch["version"], 1)
		with self.assertRaises(_native_frappe.ValidationError):
			self.save()
		with self.assertRaises(_native_frappe.TimestampMismatchError):
			self.service.save_opening(dict(batch, rows=self.rows("200")), expected_version=0)
		self.assertEqual(self.counts(), before)

	def test_new_activation_creates_one_native_opening_and_replay_posts_nothing(self):
		batch = self.save()["batch"]
		result = self.activate(batch)
		self.assertEqual(result["batch"]["status"], "Active")
		self.assertEqual(len(result["receipt"]["source_names"]), 1)
		name = result["receipt"]["source_names"][0]
		journal = _native_frappe.get_doc("Journal Entry", name)
		self.assertEqual((journal.docstatus, journal.voucher_type, journal.is_opening), (1, "Opening Entry", "Yes"))
		self.assertGreater(_native_frappe.db.count("GL Entry", {"voucher_no": name, "is_opening": "Yes", "is_cancelled": 0}), 0)
		before = self.counts()
		replay = self.service.activate_opening(batch["name"], result["receipt"]["fingerprint"])
		self.assertEqual(replay["receipt"], result["receipt"])
		self.assertEqual(self.counts(), before)
		self.assertFalse(journal.custom_china_voucher_number)
		with self.assertRaises(_native_frappe.ValidationError):
			journal.cancel()

	def test_takeover_retains_original_remarks_and_does_not_change_any_gl_or_source(self):
		journal = self.historical()
		fields = ["name", "modified", "account", "account_currency", "debit", "credit", "party_type", "party", "is_opening", "is_cancelled"]
		before = exact_snapshot = _native_frappe.get_all("GL Entry", filters={"voucher_no": journal.name}, fields=fields, order_by="name")
		before_count = self.counts()
		evidence = self.service._source_evidence(journal.name, self.company, self.opening_date)
		self.service._lock_company(self.company)
		self.assertEqual(evidence["hash"], self.service._source_evidence(journal.name, self.company, self.opening_date, lock=True)["hash"])
		batch = self.save(rows=evidence["rows"], mode="Takeover", source_names=[journal.name])["batch"]
		result = self.activate(batch)
		self.assertEqual(result["receipt"]["created_journal_entries"], [])
		self.assertEqual(self.counts(), before_count)
		self.assertEqual(_native_frappe.get_all("GL Entry", filters={"voucher_no": journal.name}, fields=fields, order_by="name"), exact_snapshot)
		self.assertEqual(result["batch"]["rows"][0]["remarks"], "原始摘要 未建档往来名称")

	def test_adjustment_holds_full_target_posts_only_delta_and_zero_delta_has_no_je(self):
		active = self.activate(self.save()["batch"])["batch"]
		adjustment = self.service.start_adjustment(active["name"], active["version"])["batch"]
		self.assertEqual(len(adjustment["rows"]), 2)
		updated = self.service.save_opening({"name": adjustment["name"], "company": self.company, "opening_date": self.opening_date, "mode": "Adjustment", "rows": self.rows("150.00")}, adjustment["version"])["batch"]
		preview = self.service.preview_opening(updated["name"])
		self.assertEqual(preview["errors"], [], preview)
		self.assertEqual(preview["totals"]["debit"], "150.00")
		self.assertEqual(core(self).totals(preview["delta_rows"])["debit"], "50.00")
		result = self.service.activate_opening(updated["name"], preview["fingerprint"])
		journal = _native_frappe.get_doc("Journal Entry", result["receipt"]["source_names"][0])
		self.assertEqual(Decimal(str(journal.total_debit)), Decimal("50.00"))
		before = self.counts()
		zero = self.service.start_adjustment(result["batch"]["name"], result["batch"]["version"])["batch"]
		zero_result = self.activate(zero)
		self.assertTrue(zero_result["receipt"]["zero_delta"])
		self.assertEqual(zero_result["receipt"]["source_names"], [])
		self.assertEqual(self.counts(), before)

	def test_failed_native_submit_rolls_back_insert_gl_snapshot_and_active_status(self):
		from unittest.mock import patch

		batch = self.save()["batch"]
		preview = self.service.preview_opening(batch["name"])
		before = self.counts()
		with patch("china_finance.services.voucher._complete_voucher_workflow", side_effect=_native_frappe.ValidationError("denied workflow")):
			with self.assertRaisesRegex(_native_frappe.ValidationError, "denied workflow"):
				self.service.activate_opening(batch["name"], preview["fingerprint"])
		self.assertEqual(self.counts(), before)
		self.assertEqual(_native_frappe.db.get_value(self.service.BATCH, batch["name"], "status"), "Draft")
		self.assertIsNone(self.service._native.get())

	def test_freeze_remains_in_force_and_activation_cannot_unfreeze(self):
		batch = self.save()["batch"]
		preview = self.service.preview_opening(batch["name"])
		_native_frappe.db.set_value("Company", self.company, "accounts_frozen_till_date", self.opening_date)
		before = self.counts()
		with self.assertRaises(_native_frappe.ValidationError):
			self.service.activate_opening(batch["name"], preview["fingerprint"])
		self.assertEqual(self.counts(), before)
		self.assertEqual(str(_native_frappe.db.get_value("Company", self.company, "accounts_frozen_till_date")), self.opening_date)

	def test_source_drift_is_rejected_on_idempotent_retry(self):
		result = self.activate(self.save()["batch"])
		name = result["receipt"]["source_names"][0]
		gl = _native_frappe.db.get_value("GL Entry", {"voucher_no": name}, "name")
		_native_frappe.db.set_value("GL Entry", gl, "is_opening", "No", update_modified=False)
		before = self.counts()
		with self.assertRaises(_native_frappe.ValidationError):
			self.service.activate_opening(result["batch"]["name"], result["receipt"]["fingerprint"])
		self.assertEqual(self.counts(), before)

	def test_guest_cannot_read_save_or_activate_and_direct_rest_fields_are_protected(self):
		batch = self.save()["batch"]
		before = self.counts()
		_native_frappe.set_user("Guest")
		for call in (
			lambda: self.service.get_opening(self.company),
			lambda: self.service.save_opening({"name": batch["name"], "company": self.company, "opening_date": self.opening_date, "rows": self.rows()}, 1),
			lambda: self.service.activate_opening(batch["name"], "forged"),
		):
			with self.assertRaises(_native_frappe.PermissionError):
				call()
		_native_frappe.set_user("Administrator")
		doc = _native_frappe.get_doc(self.service.BATCH, batch["name"])
		doc.status = "Active"
		with self.assertRaises(_native_frappe.PermissionError):
			doc.save()
		self.assertEqual(self.counts(), before)

	def test_foreign_currency_native_activation_preserves_account_and_base_amounts(self):
		foreign = self._account("Opening USD Bank", self.bank, account_type="Bank", currency="USD")
		rows = self.rows("700.00")
		rows[0].update(account=foreign, account_currency="USD", debit="100.00", base_debit="700.00", exchange_rate="7")
		result = self.activate(self.save(rows=rows)["batch"])
		journal = _native_frappe.get_doc("Journal Entry", result["receipt"]["source_names"][0])
		row = next(row for row in journal.accounts if row.account == foreign)
		self.assertEqual((Decimal(str(row.debit_in_account_currency)), Decimal(str(row.debit)), row.account_currency), (Decimal("100"), Decimal("700"), "USD"))
		self.assertEqual(row.exchange_rate, 7)

	def test_invalid_native_account_party_dimension_and_unbalanced_rows_stay_drafts(self):
		rows = self.rows()
		rows[0].update(account=self.expense_account, party_type="Customer", party="not-a-real-customer", dimensions_json='{"cost_center":"not-a-real-cost-center"}')
		rows[1]["credit"] = rows[1]["base_credit"] = "99.99"
		batch = self.save(rows=rows)["batch"]
		preview = self.service.preview_opening(batch["name"])
		self.assertTrue({"account", "party", "cost_center", "totals"}.issubset({error["field"] for error in preview["errors"]}), preview)
		before = self.counts()
		with self.assertRaises(_native_frappe.ValidationError):
			self.service.activate_opening(batch["name"], preview["fingerprint"])
		self.assertEqual(self.counts(), before)

	def test_import_preview_never_saves_batch_je_or_gl(self):
		before = self.counts()
		batch_count = _native_frappe.db.count(self.service.BATCH, {"company": self.company})
		text = "科目\t借方\t贷方\t摘要\n" + self.bank + "\t100.00\t0\t原摘要\n" + self.equity + "\t0\t100.00\t资本来源"
		preview = self.service.preview_import(self.company, text=text)
		self.assertEqual(preview["errors"], [])
		self.assertEqual(len(preview["rows"]), 2)
		self.assertEqual(self.counts(), before)
		self.assertEqual(_native_frappe.db.count(self.service.BATCH, {"company": self.company}), batch_count)

	def test_valid_master_change_invalidates_saved_preview_token(self):
		batch = self.save()["batch"]
		preview = self.service.preview_opening(batch["name"])
		_native_frappe.db.set_value("Account", self.bank, "account_name", "仍然有效的新科目名称")
		before = self.counts()
		with self.assertRaises(_native_frappe.TimestampMismatchError):
			self.service.activate_opening(batch["name"], preview["fingerprint"])
		self.assertEqual(self.counts(), before)

	def test_opening_before_finance_activation_still_retains_technical_snapshot(self):
		_native_frappe.db.set_value("China Finance Settings", self.company, "activation_date", "2027-01-01")
		_native_frappe.clear_document_cache("China Finance Settings", self.company)
		result = self.activate(self.save()["batch"])
		name = result["receipt"]["source_names"][0]
		self.assertEqual(_native_frappe.db.count("China Accounting Voucher", {"company": self.company, "source_name": name, "source_event": "Posting", "docstatus": 1}), 1)

	def test_disabled_finance_never_activates_without_technical_snapshot(self):
		_native_frappe.db.set_value("China Finance Settings", self.company, "enabled", 0)
		_native_frappe.clear_document_cache("China Finance Settings", self.company)
		batch = self.save()["batch"]
		preview = self.service.preview_opening(batch["name"])
		before = self.counts()
		with self.assertRaises(_native_frappe.ValidationError):
			self.service.activate_opening(batch["name"], preview["fingerprint"])
		self.assertEqual(self.counts(), before)

	def test_native_same_named_closed_document_and_period_disabled_change_tokens(self):
		period = _native_frappe.get_doc({
			"doctype": "Accounting Period", "period_name": "_Opening Period " + _native_frappe.generate_hash(length=10),
			"company": self.company, "start_date": "2026-01-01", "end_date": "2026-12-31",
			"closed_documents": [{"document_type": "Journal Entry", "closed": 0}],
		}).insert()
		batch = self.save()["batch"]
		first = self.service.preview_opening(batch["name"])
		child_name = period.closed_documents[0].name
		_native_frappe.db.set_value("Closed Document", child_name, "closed", 1, update_modified=False)
		second = self.service.preview_opening(batch["name"])
		self.assertNotEqual(first["fingerprint"], second["fingerprint"])
		self.assertEqual(_native_frappe.get_doc("Accounting Period", period.name).closed_documents[0].name, child_name)
		_native_frappe.db.set_value("Closed Document", child_name, "closed", 0, update_modified=False)
		third = self.service.preview_opening(batch["name"])
		_native_frappe.db.set_value("Accounting Period", period.name, "disabled", 1, update_modified=False)
		fourth = self.service.preview_opening(batch["name"])
		self.assertNotEqual(third["fingerprint"], fourth["fingerprint"])
		with self.assertRaises(_native_frappe.TimestampMismatchError):
			self.service.activate_opening(batch["name"], third["fingerprint"])

	def test_more_than_one_hundred_rows_submit_atomically_without_background_queue(self):
		from unittest.mock import patch

		rows = [dict(row, row_id=row["row_id"] + str(index)) for index in range(51) for row in self.rows("1.00")]
		batch = self.save(rows=rows)["batch"]
		with patch("erpnext.accounts.doctype.journal_entry.journal_entry.queue_submission", side_effect=AssertionError("Opening activation must never enqueue accounting posting")) as queued:
			result = self.activate(batch)
		queued.assert_not_called()
		journal = _native_frappe.get_doc("Journal Entry", result["receipt"]["source_names"][0])
		self.assertEqual(journal.docstatus, 1)
		self.assertEqual(len(journal.accounts), 102)
		self.assertEqual(Decimal(str(journal.total_debit)), Decimal("51"))

	def test_company_and_batch_lock_order_is_explicit(self):
		from unittest.mock import patch

		batch = self.save()["batch"]
		events = []
		original = _native_frappe.get_doc

		def read(*args, **kwargs):
			if kwargs.get("for_update") and args and args[0] == self.service.BATCH:
				events.append("batch")
			return original(*args, **kwargs)

		with patch.object(self.service, "_lock_company", side_effect=lambda company: events.append("company")), patch("frappe.get_doc", side_effect=read):
			self.service._get_batch(batch["name"], write=True, lock=True)
		self.assertEqual(events, ["company", "batch"])


@unittest.skipUnless(_native_frappe, "Requires isolated native MariaDB connections")
class TestOpeningCurrentReadConcurrency(unittest.TestCase):
	def setUp(self):
		if _native_frappe.local.site not in ("cash-flow-list-qa.localhost", "dev.localhost") or _native_frappe.db.db_type != "mariadb":
			self.skipTest("Committed concurrent fixtures are restricted to local MariaDB QA sites")

	def scenario(self, change):
		import subprocess
		import sys

		code = """
import sys
import frappe
frappe.init(site=sys.argv[1], sites_path=sys.argv[2])
frappe.connect()
frappe.flags.in_test = True
try:
    from china_finance.tests.test_opening_balance import run_opening_current_read_scenario
    run_opening_current_read_scenario(sys.argv[3])
finally:
    frappe.db.rollback()
    frappe.destroy()
"""
		try:
			result = subprocess.run([sys.executable, "-c", code, _native_frappe.local.site, str(_native_frappe.local.sites_path), change], capture_output=True, text=True, timeout=45)
		finally:
			# The parent test owns no business writes. Child commits deliberately
			# change then remove Account fixtures; discard the parent's old RR
			# read view before later native nested-set fixture construction.
			_native_frappe.db.rollback()
		self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

	def test_account_change_committed_while_activation_waits_rejects_original_preview(self):
		self.scenario("Account")

	def test_closed_document_change_while_activation_waits_rejects_original_preview(self):
		self.scenario("Closed Document")

	def test_native_cancel_waits_before_je_lock_and_is_rejected_after_takeover(self):
		self.scenario("Native Cancel")


def run_opening_current_read_scenario(change):
	"""Two real transactions: ordinary RR read, wait, then changed master evidence.

	Only a randomly named isolated company and its fixture chain are committed;
	activation is expected to fail before any accounting write and is rolled back.
	"""
	import threading
	from concurrent.futures import ThreadPoolExecutor
	from unittest.mock import patch

	import frappe

	from china_finance.services import opening_balance as service
	from china_finance.services.month_end import lock_company
	from china_finance.tests.test_voucher_numbering_concurrency import _connection_action

	if frappe.local.site not in ("cash-flow-list-qa.localhost", "dev.localhost") or frappe.db.db_type != "mariadb":
		raise AssertionError("Opening concurrency fixtures are restricted to local QA MariaDB")
	frappe.set_user("Administrator")
	frappe.db.sql("SET SESSION innodb_lock_wait_timeout=5")
	assert frappe.db.sql("SELECT @@tx_isolation")[0][0] == "REPEATABLE-READ"
	prefix = "_OPEN_CONC_" + frappe.generate_hash(length=12)
	company = prefix + " Company"
	period_name = prefix + " Period"
	accounts = [prefix + " Bank", prefix + " Equity"]
	waiting = threading.Event()
	site, sites_path = frappe.local.site, str(frappe.local.sites_path)
	try:
		frappe.get_doc({"doctype": "Company", "name": company, "company_name": company, "abbr": frappe.generate_hash(length=5), "default_currency": "CNY"}).db_insert()
		for name, root_type in zip(accounts, ("Asset", "Equity")):
			frappe.get_doc({"doctype": "Account", "name": name, "account_name": name, "company": company, "is_group": 0, "disabled": 0, "root_type": root_type, "report_type": "Balance Sheet", "account_currency": "CNY"}).db_insert()
		period = frappe.get_doc({"doctype": "Accounting Period", "name": period_name, "period_name": period_name, "company": company, "start_date": "2026-01-01", "end_date": "2026-12-31", "disabled": 0, "closed_documents": [{"document_type": "Journal Entry", "closed": 0}]})
		period.db_insert()
		child = period.closed_documents[0]
		child.parent, child.parenttype, child.parentfield = period_name, "Accounting Period", "closed_documents"
		child.db_insert()
		payload = {"company": company, "opening_date": "2026-01-01", "mode": "New", "rows": [
			{"row_id": "bank", "account": accounts[0], "account_currency": "CNY", "exchange_rate": 1, "debit": 100, "base_debit": 100},
			{"row_id": "equity", "account": accounts[1], "account_currency": "CNY", "exchange_rate": 1, "credit": 100, "base_credit": 100},
		]}
		if change == "Native Cancel":
			source = frappe.get_doc({"doctype": "Journal Entry", "name": prefix + " JE", "company": company, "docstatus": 1, "posting_date": "2026-01-01", "voucher_type": "Opening Entry", "is_opening": "Yes", "accounts": [
				{"account": row["account"], "account_currency": "CNY", "exchange_rate": 1, "debit": row.get("base_debit", 0), "credit": row.get("base_credit", 0), "debit_in_account_currency": row.get("debit", 0), "credit_in_account_currency": row.get("credit", 0), "user_remark": "Opening source"} for row in payload["rows"]
			]})
			source.db_insert()
			for row in source.accounts:
				row.parent, row.parenttype, row.parentfield = source.name, "Journal Entry", "accounts"
				row.db_insert()
				frappe.get_doc({"doctype": "GL Entry", "name": prefix + " GL " + str(row.idx), "company": company, "posting_date": "2026-01-01", "voucher_type": "Journal Entry", "voucher_no": source.name, "voucher_detail_no": row.name, "is_opening": "Yes", "is_cancelled": 0, "account": row.account, "account_currency": "CNY", "debit": row.debit, "credit": row.credit, "debit_in_account_currency": row.debit_in_account_currency, "credit_in_account_currency": row.credit_in_account_currency}).db_insert()
			payload.update(mode="Takeover", source_names=[source.name], rows=service._source_evidence(source.name, company, "2026-01-01")["rows"])
		batch = service.save_opening(payload)["batch"]
		preview = service.preview_opening(batch["name"])
		assert not preview["errors"], preview
		frappe.db.commit()
		lock_company(company)

		def reader(retrying=False):
			def gate(value):
				waiting.set()
				lock_company(value)
			if change == "Native Cancel":
				journal = frappe.get_doc("Journal Entry", source.name)
				with patch("china_finance.services.month_end.lock_company", side_effect=gate):
					try:
						journal.cancel()
					except frappe.ValidationError as exc:
						assert "期初余额" in str(exc), exc
						return "opening-protected"
					raise AssertionError("Native cancel bypassed activated opening protection")
			with patch.object(service, "_lock_company", side_effect=gate):
				try:
					service.activate_opening(batch["name"], preview["fingerprint"])
				except frappe.TimestampMismatchError:
					return "stale-preview"
				raise AssertionError("Activation accepted a master change committed after its initial RR read")

		with ThreadPoolExecutor(max_workers=1) as executor:
			future = executor.submit(_connection_action, site, sites_path, reader)
			try:
				if not waiting.wait(10):
					if future.done():
						future.result()
					raise AssertionError("Activation did not reach its company gate")
				if change == "Native Cancel":
					result = service.activate_opening(batch["name"], preview["fingerprint"])
					assert result["batch"]["status"] == "Active", result
				elif change == "Account":
					frappe.db.set_value("Account", accounts[0], "account_name", "Changed while waiting", update_modified=False)
				else:
					frappe.db.set_value("Closed Document", child.name, "closed", 1, update_modified=False)
				frappe.db.commit()
				assert future.result(timeout=15) == ("opening-protected" if change == "Native Cancel" else "stale-preview")
			finally:
				frappe.db.rollback()
		assert frappe.db.count("Journal Entry", {"company": company}) == (1 if change == "Native Cancel" else 0)
		assert frappe.db.count("GL Entry", {"company": company}) == (2 if change == "Native Cancel" else 0)
		if change == "Native Cancel":
			assert frappe.db.get_value("Journal Entry", source.name, "docstatus") == 1
			assert frappe.db.count("GL Entry", {"company": company, "is_cancelled": 1}) == 0
	finally:
		frappe.db.rollback()
		for name in frappe.get_all(service.BATCH, filters={"company": company}, pluck="name"):
			frappe.db.delete(service.LINE, {"parent": name})
		frappe.db.delete(service.BATCH, {"company": company})
		frappe.db.delete("Closed Document", {"parent": period_name})
		frappe.db.delete("Accounting Period", {"name": period_name})
		for name in frappe.get_all("Journal Entry", filters={"company": company}, pluck="name"):
			frappe.db.delete("Journal Entry Account", {"parent": name})
		frappe.db.delete("GL Entry", {"company": company})
		frappe.db.delete("Journal Entry", {"company": company})
		frappe.db.delete("Account", {"company": company})
		frappe.db.delete("Company", {"name": company})
		frappe.db.commit()


if __name__ == "__main__":
	unittest.main()
